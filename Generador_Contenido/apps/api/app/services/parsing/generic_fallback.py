"""Best-effort extraction used when a document arrives in the "wrong" format
for its role (e.g. a sílabo uploaded as PDF instead of DOCX, or a guía
uploaded as DOCX instead of PDF). The primary parsers (docx_parser.py,
pdf_parser.py) rely on this institution's real table/page layout for high
-confidence extraction; this fallback only guarantees week numbers and
competency codes can still be found, with a warning to the teacher that
accuracy is reduced.
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
import pdfplumber

from app.schemas.extraction import GuiaExtraction, SilaboExtraction, SilaboWeek
from app.services.parsing.config import WEEK_HEADER_PATTERN, extract_codes


def extract_raw_text(path: str | Path, ext: str) -> str:
    if ext == ".docx":
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if ext == ".pdf":
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    raise ValueError(f"unsupported extension: {ext}")


def _split_by_week(text: str) -> list[tuple[int, str]]:
    matches = list(WEEK_HEADER_PATTERN.finditer(text))
    chunks: list[tuple[int, str]] = []
    for i, m in enumerate(matches):
        start = m.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunks.append((int(m.group(1)), text[start:end]))
    return chunks


def fallback_silabo_extraction(text: str) -> SilaboExtraction:
    chunks = _split_by_week(text)
    weeks = []
    for week_number, chunk in chunks:
        lines = [l.strip() for l in chunk.splitlines() if l.strip()]
        topic = lines[1] if len(lines) > 1 else ""
        weeks.append(
            SilaboWeek(
                week_number=week_number,
                topic=topic[:200],
                competency_codes=extract_codes(chunk),
                confidence=0.5,
            )
        )
    warnings = [
        "Extracción limitada: este documento no tiene el formato esperado para un sílabo "
        "(tabla DOCX). Se detectaron semanas y códigos de competencia por texto, pero revise "
        "manualmente los demás campos."
    ]
    return SilaboExtraction(weeks=weeks, warnings=warnings)


def fallback_guia_extraction(text: str) -> GuiaExtraction:
    chunks = _split_by_week(text)
    weeks_detected = [wn for wn, _ in chunks]
    codes_by_week = {wn: extract_codes(chunk) for wn, chunk in chunks}
    warnings = [
        "Extracción limitada: este documento no tiene el formato esperado para una guía "
        "didáctica (PDF paginado). Se detectaron semanas y códigos de competencia por texto, "
        "pero no fue posible ubicar secciones ni páginas."
    ]
    return GuiaExtraction(
        weeks_detected=sorted(set(weeks_detected)),
        competency_codes_by_week=codes_by_week,
        warnings=warnings,
    )
