"""Renders a GeneratedResource into a .pptx, from the same structured
`block.payload` used to render HTML — not by parsing rendered_html — so
interactive elements degrade to readable static slides (e.g. a quiz becomes
one slide per question with the correct option marked) instead of being lost.
"""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

from app.schemas.generation import GeneratedResource, ResourceBlockType


def _title_slide(prs: Presentation, resource: GeneratedResource) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = resource.title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = resource.topic


def _bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title or "Recurso"
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, bullet in enumerate(bullets or [""]):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)


def _quiz_bullets(item: dict) -> list[str]:
    options = item.get("options", [])
    correct = item.get("correct_index")
    return [f"{'✓ ' if i == correct else '• '}{opt}" for i, opt in enumerate(options)]


def _word_search_slide(prs: Presentation, block) -> None:
    words = block.payload.get("words", [])
    title = block.title or "Sopa de Letras"

    from app.mcp_server.tools.word_search import _generate_grid
    grid, placed = _generate_grid(words, size=block.payload.get("grid_size", 12))

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_box = slide.shapes.add_textbox(Pt(30), Pt(20), Pt(650), Pt(40))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True

    # Grid Table
    rows = len(grid)
    cols = len(grid[0]) if rows > 0 else 0
    if rows > 0 and cols > 0:
        table_shape = slide.shapes.add_table(rows, cols, Pt(30), Pt(70), Pt(420), Pt(420))
        table = table_shape.table

        for r in range(rows):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = grid[r][c]
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(10)
                p.font.bold = True
                p.alignment = 1

    # Words List
    word_box = slide.shapes.add_textbox(Pt(470), Pt(70), Pt(240), Pt(420))
    wtf = word_box.text_frame
    wp = wtf.paragraphs[0]
    wp.text = f"Palabras a buscar ({len(placed)}):"
    wp.font.size = Pt(14)
    wp.font.bold = True

    for w in placed:
        p = wtf.add_paragraph()
        p.text = f"• {w}"
        p.font.size = Pt(11)


def _flashcard_slide(prs: Presentation, card: dict, idx: int, total: int, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    header_box = slide.shapes.add_textbox(Pt(30), Pt(15), Pt(650), Pt(35))
    tf = header_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"{title} — Tarjeta {idx+1}/{total}"
    p.font.size = Pt(18)
    p.font.bold = True

    # Front Box (Blue)
    front_shape = slide.shapes.add_shape(1, Pt(30), Pt(65), Pt(320), Pt(410))
    front_shape.fill.solid()
    front_shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
    front_shape.line.color.rgb = RGBColor(59, 130, 246)
    front_shape.line.width = Pt(2)

    ftf = front_shape.text_frame
    ftf.word_wrap = True
    p0 = ftf.paragraphs[0]
    p0.text = f"CONCEPTO / PREGUNTA #{idx+1}"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(29, 78, 216)

    p1 = ftf.add_paragraph()
    p1.text = f"\n{card.get('front', '')}"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(30, 41, 59)

    # Back Box (Green)
    back_shape = slide.shapes.add_shape(1, Pt(370), Pt(65), Pt(320), Pt(410))
    back_shape.fill.solid()
    back_shape.fill.fore_color.rgb = RGBColor(240, 253, 244)
    back_shape.line.color.rgb = RGBColor(34, 197, 94)
    back_shape.line.width = Pt(2)

    btf = back_shape.text_frame
    btf.word_wrap = True
    bp0 = btf.paragraphs[0]
    bp0.text = "EXPLICACIÓN / RESPUESTA"
    bp0.font.size = Pt(11)
    bp0.font.bold = True
    bp0.font.color.rgb = RGBColor(21, 128, 61)

    bp1 = btf.add_paragraph()
    bp1.text = f"\n{card.get('back', '')}"
    bp1.font.size = Pt(14)
    bp1.font.color.rgb = RGBColor(20, 83, 45)


def _study_guide_slide(prs: Presentation, block) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = block.title or "Guía de Estudio"
    header_box = slide.shapes.add_textbox(Pt(30), Pt(15), Pt(650), Pt(35))
    tf = header_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"📌 {title}"
    p.font.size = Pt(20)
    p.font.bold = True

    sections = block.payload.get("sections", [])
    num_sec = max(1, min(len(sections), 3))
    w_per_col = 660 / num_sec

    for i, sec in enumerate(sections[:3]):
        left_pos = Pt(30 + i * w_per_col)
        sec_shape = slide.shapes.add_shape(1, left_pos, Pt(65), Pt(w_per_col - 15), Pt(410))
        sec_shape.fill.solid()
        sec_shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        sec_shape.line.color.rgb = RGBColor(203, 213, 225)
        sec_shape.line.width = Pt(1.5)

        stf = sec_shape.text_frame
        stf.word_wrap = True
        p0 = stf.paragraphs[0]
        p0.text = sec.get("heading", "")
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(15, 23, 42)

        for b in sec.get("bullets", []):
            bp = stf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.size = Pt(11)
            bp.font.color.rgb = RGBColor(51, 65, 85)

        if sec.get("key_takeaway"):
            kp = stf.add_paragraph()
            kp.text = f"\n💡 Idea Clave:\n{sec.get('key_takeaway')}"
            kp.font.size = Pt(10)
            kp.font.bold = True
            kp.font.color.rgb = RGBColor(146, 64, 14)


def _diagram_slide(prs: Presentation, block) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = block.title or "Esquema / Mapa Conceptual"
    header_box = slide.shapes.add_textbox(Pt(30), Pt(15), Pt(650), Pt(35))
    tf = header_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"🗺️ {title}"
    p.font.size = Pt(20)
    p.font.bold = True

    nodes = block.payload.get("nodes", [])
    if not nodes:
        return

    n_nodes = len(nodes)
    cols = min(3, max(1, n_nodes))
    w_card = Pt(200)
    h_card = Pt(90)

    for i, node in enumerate(nodes[:6]):
        row = i // cols
        col = i % cols
        left = Pt(30 + col * 225)
        top = Pt(65 + row * 115)

        shape = slide.shapes.add_shape(1, left, top, w_card, h_card)
        shape.fill.solid()
        ntype = node.get("node_type", "concept")

        if ntype == "concept":
            shape.fill.fore_color.rgb = RGBColor(238, 246, 255)
            shape.line.color.rgb = RGBColor(59, 130, 246)
        elif ntype == "process":
            shape.fill.fore_color.rgb = RGBColor(240, 253, 244)
            shape.line.color.rgb = RGBColor(34, 197, 94)
        elif ntype == "example":
            shape.fill.fore_color.rgb = RGBColor(254, 243, 199)
            shape.line.color.rgb = RGBColor(245, 158, 11)
        else:
            shape.fill.fore_color.rgb = RGBColor(243, 232, 255)
            shape.line.color.rgb = RGBColor(168, 85, 247)

        shape.line.width = Pt(1.5)
        stf = shape.text_frame
        stf.word_wrap = True
        p0 = stf.paragraphs[0]
        p0.text = node.get("label", "")
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(15, 23, 42)


def export_pptx(resource: GeneratedResource) -> bytes:
    prs = Presentation()
    _title_slide(prs, resource)

    for block in resource.blocks:
        if block.type == ResourceBlockType.interactive_activity:
            for item in block.payload.get("items", []):
                _bullet_slide(prs, item.get("question", block.title), _quiz_bullets(item))
        elif block.type == ResourceBlockType.word_search:
            _word_search_slide(prs, block)
        elif block.type == ResourceBlockType.flashcards:
            cards = block.payload.get("cards", [])
            for idx, c in enumerate(cards):
                _flashcard_slide(prs, c, idx, len(cards), block.title or "Tarjetas de Estudio")
        elif block.type == ResourceBlockType.study_guide:
            _study_guide_slide(prs, block)
        elif block.type == ResourceBlockType.diagram:
            _diagram_slide(prs, block)
        else:
            _bullet_slide(prs, block.title or "Recurso", [block.title or ""])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
